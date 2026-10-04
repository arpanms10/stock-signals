"""Turn a Kite holdings download into config/holdings.csv.

Drop the file Kite gives you into config/ and run this. Both exports work:

    Kite web   Holdings -> Download   .csv   Instrument, Qty., Avg. cost, ...
    Console    Portfolio -> Holdings  .xlsx  a client-details preamble, then
                                             Symbol, Quantity Available,
                                             Quantity Long Term, Average Price

The header row is found wherever it sits, so the Console preamble and its
Mutual Funds sheet are skipped rather than misread. Bonds, REITs, InvITs and
ETFs are left out, and listed, because the framework ranks equities only.

An existing holdings.csv is not thrown away: it is copied to holdings.csv.bak,
and the bucket, notes and purchase_date you set by hand are carried over for
every symbol you still hold. Kite knows your quantities; only you know why you
own something.

    PYTHONPATH=. .venv/bin/python import_kite_holdings.py            # newest Kite file in config/
    PYTHONPATH=. .venv/bin/python import_kite_holdings.py path/to/holdings.xlsx
    PYTHONPATH=. .venv/bin/python import_kite_holdings.py --satellite CPPLUS SANSERA
"""
from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from pathlib import Path

import instruments
import portfolio as pf

CONFIG = Path(__file__).parent / "config"

# Files in config/ that belong to the framework, never a broker export.
OWN_FILES = {"holdings.csv", "holdings.csv.bak", "holdings.example.csv",
             "fundamentals.csv", "fundamentals.example.csv", "watching.csv",
             "watchlist.csv", "equities.csv", "etfs.csv"}

# Console splits the position across several columns; the holding is their sum.
# Pledged shares are still yours -- leaving them out would understate the book.
QUANTITY_PARTS = ("quantity available", "quantity discrepant",
                  "quantity pledged (margin)", "quantity pledged (loan)")

LT_ALIASES = pf.COLUMN_ALIASES["lt_quantity"] + ("quantity long term",)

# Kite shows other equity series as SYMBOL-BE etc. The framework keys on the
# bare NSE symbol, so the suffix has to go or the stock silently never matches.
EQUITY_SERIES = re.compile(r"-(BE|BZ|BL|BT|SM|ST)$")

# Any other short suffix is a non-equity series: -RR REIT, -IV InvIT, -NE/-N1
# bonds, -GB gold bonds, -F ETFs... A 1-2 character limit keeps real symbols
# with a hyphen, like BAJAJ-AUTO, out of it.
OTHER_SERIES = re.compile(r"-([A-Z0-9]{1,2})$")
SERIES_NAMES = {"RR": "REIT", "IV": "InvIT", "GB": "gold bond"}

# ISIN characters 8-9 are the security type; 01 is equity shares. A company's
# bonds share its INE prefix (INE530B08094 is a debenture), so the prefix alone
# cannot tell a stock from its debt.
EQUITY_SHARE_TYPE = "01"

# Hand-maintained columns that survive a re-import.
KEPT = ("bucket", "notes", "purchase_date")


def clean_symbol(raw) -> str:
    sym = str(raw or "").strip().strip('"').upper()
    return EQUITY_SERIES.sub("", sym)


def non_equity(sym: str, isin: str = "", sector: str = "") -> str | None:
    """Why this line is not a stock the framework can rank, or None if it is.

    The framework ranks and checks NSE equities only. A bond or REIT in
    holdings.csv would have no prices or filings behind it, yet still count
    towards the size of the book.
    """
    if sector.strip().lower() == "debt":
        return "debt"
    isin = isin.strip().upper()
    kind = instruments.classify(isin) if isin else ""
    if kind == "fund":
        return "ETF / fund"
    m = OTHER_SERIES.search(sym)
    if m:
        return SERIES_NAMES.get(m.group(1), f"non-equity series -{m.group(1)}")
    if isin:
        if kind != "equity" or isin[7:9] != EQUITY_SHARE_TYPE:
            return f"not an equity share (ISIN {isin})"
    return None


def _norm(h) -> str:
    return str(h or "").strip().strip('"').lower()


def _find(header: list[str], aliases) -> int | None:
    for a in aliases:
        if a in header:
            return header.index(a)
    return None


def locate_header(grid: list[list]) -> tuple[int, dict] | None:
    """Index of the header row and the column positions it implies."""
    for i, row in enumerate(grid):
        header = [_norm(c) for c in row]
        sym = _find(header, pf.COLUMN_ALIASES["symbol"])
        if sym is None:
            continue
        parts = [header.index(p) for p in QUANTITY_PARTS if p in header]
        qty = _find(header, pf.COLUMN_ALIASES["quantity"])
        if not parts and qty is None:
            continue
        return i, {
            "symbol": sym,
            "quantity": parts or [qty],
            "avg_price": _find(header, pf.COLUMN_ALIASES["avg_price"]),
            "lt_quantity": _find(header, LT_ALIASES),
            "st_quantity": _find(header, pf.COLUMN_ALIASES["st_quantity"]),
            "isin": _find(header, ("isin",)),
            "sector": _find(header, ("sector",)),
        }
    return None


def _cell(row: list, idx: int | None) -> float | None:
    if idx is None or idx >= len(row):
        return None
    return pf._to_float(row[idx])


def _text(row: list, idx: int | None) -> str:
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def convert(grid: list[list]) -> tuple[list[dict], list[tuple[str, str]]]:
    """Rows of a Kite export -> (holdings.csv rows, skipped (symbol, reason)).

    Both empty if no header row is found.
    """
    found = locate_header(grid)
    if not found:
        return [], []
    start, cols = found
    merged: dict[str, dict] = {}
    skipped: list[tuple[str, str]] = []
    for row in grid[start + 1:]:
        if cols["symbol"] >= len(row):
            continue
        sym = clean_symbol(row[cols["symbol"]])
        # A blank symbol ends the table (Console puts totals / notes below it).
        if not sym:
            if merged:
                break
            continue
        qty = sum(_cell(row, i) or 0.0 for i in cols["quantity"])
        if qty <= 0:
            continue
        why = non_equity(sym, _text(row, cols["isin"]), _text(row, cols["sector"]))
        if why:
            skipped.append((sym, why))
            continue
        avg = _cell(row, cols["avg_price"]) or 0.0
        lt = _cell(row, cols["lt_quantity"])
        st = _cell(row, cols["st_quantity"])
        if lt is not None and st is None:
            st = max(qty - lt, 0.0)
        elif st is not None and lt is None:
            lt = max(qty - st, 0.0)

        # The same stock can appear twice (an EQ line and a BE line).
        # Combine into one position at the blended cost.
        if sym in merged:
            m = merged[sym]
            total = m["quantity"] + qty
            m["avg_price"] = (m["avg_price"] * m["quantity"] + avg * qty) / total
            m["quantity"] = total
            if lt is not None:
                m["lt_quantity"] = (m["lt_quantity"] or 0.0) + lt
                m["st_quantity"] = (m["st_quantity"] or 0.0) + st
            continue
        merged[sym] = {"symbol": sym, "quantity": qty, "lt_quantity": lt,
                       "st_quantity": st, "avg_price": avg}
    return list(merged.values()), skipped


def read_grid(path: Path) -> list[list]:
    """Every row of the file as a list of cells, preamble included."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as fh:
            return list(csv.reader(fh))
    try:
        import openpyxl
    except ImportError:
        sys.exit("Reading .xlsx needs openpyxl:  uv pip install openpyxl\n"
                 "Or download the CSV from Kite instead.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Console's workbook has Equity and Mutual Funds sheets; only Equity counts.
    names = [n for n in wb.sheetnames if n.strip().lower() == "equity"] or wb.sheetnames
    grid: list[list] = []
    for name in names:
        ws = wb[name]
        # Console's file declares every sheet as 1x1, and read-only mode trusts
        # that and yields a single empty cell. Ignore it and read what is there.
        ws.reset_dimensions()
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if locate_header(rows):
            grid = rows
            break
    wb.close()
    return grid


def find_export(folder: Path = CONFIG) -> Path | None:
    """The most recently modified broker file in config/."""
    candidates = [p for p in folder.iterdir()
                  if p.suffix.lower() in (".csv", ".xlsx", ".xlsm")
                  and p.name.lower() not in OWN_FILES
                  and not p.name.startswith(("~$", "."))]
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    if candidates:
        return candidates[0]
    # Kite web names its download holdings.csv, so it may already be sitting
    # where the output goes. Convert it in place (the original goes to .bak).
    own = folder / "holdings.csv"
    if own.exists() and "kite" == _source_kind(own):
        return own
    return None


def _source_kind(path: Path) -> str:
    """'kite' if the file's header is Kite's rather than ours."""
    with path.open(encoding="utf-8-sig") as fh:
        for line in fh:
            if line.strip() and not line.lstrip().startswith("#"):
                return "ours" if _norm(line.split(",")[0]) == "symbol" else "kite"
    return "ours"


def mark_satellite(rows: list[dict], symbols) -> tuple[list[str], list[str]]:
    """Set bucket=satellite on the named holdings. Returns (marked, missing).

    A new holding otherwise arrives with a blank bucket, and the framework
    never suggests satellite on its own -- so a momentum buy would be treated
    as legacy and never rotated, which defeats the reason it was bought.
    """
    wanted = {clean_symbol(s) for s in symbols}
    marked = []
    for r in rows:
        if r["symbol"] in wanted:
            r["bucket"] = "satellite"
            marked.append(r["symbol"])
    return sorted(marked), sorted(wanted - set(marked))


def _fmt(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.2f}"
    return str(v)


def write_holdings(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        fh.write(pf.HOLDINGS_HELP)
        w = csv.DictWriter(fh, fieldnames=pf.HOLDINGS_COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow({c: _fmt(r.get(c)) for c in pf.HOLDINGS_COLUMNS})


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("source", nargs="?", type=Path,
                    help="Kite export (.csv or .xlsx); default: newest in config/")
    ap.add_argument("--out", type=Path, default=pf.HOLDINGS_CSV)
    ap.add_argument("--satellite", nargs="+", metavar="SYMBOL", default=[],
                    help="mark these holdings as satellite (momentum buys to be "
                         "rotated on rank); overrides any existing bucket")
    args = ap.parse_args()

    src = args.source or find_export()
    if not src or not src.exists():
        sys.exit("No Kite holdings file found. Put the .csv or .xlsx from Kite "
                 "into config/ and run again.")
    rows, skipped = convert(read_grid(src))
    if not rows and not skipped:
        sys.exit(f"{src.name}: no holdings table found (looked for a header row "
                 "with Instrument/Symbol and Qty./Quantity).")

    out: Path = args.out
    previous = {h["tradingsymbol"]: h for h in pf.read_csv_holdings(out)}
    for r in rows:
        old = previous.get(r["symbol"], {})
        for k in KEPT:
            r[k] = old.get(k, "")
    rows.sort(key=lambda r: r["symbol"])
    marked, missing = mark_satellite(rows, args.satellite)

    if out.exists():
        shutil.copy2(out, out.with_name(out.name + ".bak"))
    out.parent.mkdir(parents=True, exist_ok=True)
    write_holdings(rows, out)

    print(f"Read {src.name}: {len(rows)} positions -> {out}")
    if previous:
        now = {r["symbol"] for r in rows}
        # A skipped line is still held; it is reported below, not as sold.
        held = now | {sym for sym, _ in skipped}
        gone, new = sorted(set(previous) - held), sorted(now - set(previous))
        kept = sum(1 for r in rows if r["bucket"])
        print(f"  previous file backed up to {out.name}.bak; "
              f"bucket kept for {kept} position(s)")
        if new:
            print(f"  new:     {', '.join(new)}")
            unmarked = [s for s in new if s not in marked]
            if unmarked and not args.satellite:
                print("           a new holding with a blank bucket is NOT "
                      "satellite -- if these\n           were momentum buys, "
                      "re-run with --satellite " + " ".join(unmarked))
        if gone:
            print(f"  removed: {', '.join(gone)}")
    if marked:
        print(f"  marked satellite: {', '.join(marked)}")
    if missing:
        print(f"  !! not in this Kite file, so not marked: {', '.join(missing)}")
    if skipped:
        print(f"  skipped {len(skipped)} non-equity holding(s) -- not in holdings.csv:")
        for sym, why in skipped:
            print(f"    {sym:<16} {why}")
    no_lt = [r["symbol"] for r in rows if r["lt_quantity"] is None]
    if no_lt:
        print("  note: this export has no long/short-term split; tax figures "
              "will fall back to purchase_date (Console's .xlsx includes it)")


if __name__ == "__main__":
    main()

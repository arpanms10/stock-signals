"""Forward log of the F&O read: record each cycle's ranges, score them at expiry.

The validation (run_fno_backtest.py) says how the range WOULD have done on
history. This file says how it actually did, from ranges written down before
the outcome was known -- the only test that cannot be fitted to the past.

Writes history/fno_ranges.xlsx:

  one tab per monthly expiry ("Oct 2026 expiry"): every recorded underlying's
      spot, straddle range, the hit rate the backtest expects, the walls,
      PCR and max pain -- then, once expiry has passed, its expiry close and
      whether each level held.
  Summary: per expiry and in total, actual hit rates against expected.

Each run first fills in any expired cycle, then records the current cycle if
it is not recorded yet. The first record of a cycle is kept; --force
replaces it. The Saturday job runs it, so a cycle is normally recorded on the
first Saturday after the previous expiry, ~15-20 sessions out.

Expiry closes come from NSE's F&O bhavcopy (the underlying's close on expiry
day), the same source as the validation. It is published the evening of
expiry, so a cycle is scored on the next run after that.

    PYTHONPATH=. .venv/bin/python record_fno.py
    PYTHONPATH=. .venv/bin/python record_fno.py --symbols NIFTY RELIANCE
    PYTHONPATH=. .venv/bin/python record_fno.py --check-only
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

HISTORY = Path(__file__).parent / "history" / "fno_ranges.xlsx"
SUMMARY = "Summary"
INDICES = ["NIFTY", "BANKNIFTY", "FINNIFTY", "MIDCPNIFTY"]

RECORD_COLS = ["recorded_on", "nse_time", "expiry", "sessions_left", "symbol",
               "kind", "spot", "straddle", "range_low", "range_high",
               "expected_hit_pct", "atm_iv_pct", "put_wall", "call_wall",
               "pcr_oi", "pcr_positioning", "max_pain", "iv_straddle_pct",
               "iv_percentile", "iv_rv", "skew_pts", "term_pts",
               "results_before_expiry"]
OUTCOME_COLS = ["expiry_close", "ret_pct", "inside_range", "put_wall_held",
                "call_wall_held", "max_pain_closer", "scored_on"]
COLUMNS = RECORD_COLS + OUTCOME_COLS


def sheet_name(expiry: dt.date) -> str:
    return f"{expiry:%b %Y} expiry"


def _num(v, nd: int = 2):
    if v is None or v != v:
        return None
    return round(float(v), nd)


def record_row(v, today: dt.date) -> dict:
    """One underlying's read, as written before the outcome is known."""
    return {
        "recorded_on": today, "nse_time": v.timestamp, "expiry": v.expiry,
        "sessions_left": v.sessions, "symbol": v.symbol, "kind": v.kind,
        "spot": _num(v.spot), "straddle": _num(v.straddle),
        "range_low": _num(v.range_low), "range_high": _num(v.range_high),
        "expected_hit_pct": _num(v.range_hit_pct, 0),
        "atm_iv_pct": _num(v.atm_iv, 1),
        "put_wall": _num(v.support.strike) if v.support else None,
        "call_wall": _num(v.resistance.strike) if v.resistance else None,
        "pcr_oi": _num(v.pcr["oi"]), "pcr_positioning": v.pcr_bias,
        "max_pain": _num(v.max_pain),
        "iv_straddle_pct": _num(v.vol.iv, 1) if v.vol else None,
        "iv_percentile": _num(v.vol.iv_pct, 0) if v.vol else None,
        "iv_rv": _num(v.vol.iv_rv) if v.vol else None,
        "skew_pts": _num(v.vol.skew, 1) if v.vol else None,
        "term_pts": _num(v.vol.term, 1) if v.vol else None,
        "results_before_expiry": (", ".join(f"{d:%d-%b}" for d in v.vol.results)
                                  if v.vol and v.vol.results else None),
    }


def score_row(r: dict, close: float, today: dt.date) -> dict:
    """Outcome columns for one recorded row, given the expiry close."""
    def between(lo, hi):
        return None if lo is None or hi is None else bool(lo <= close <= hi)

    spot, mp = r.get("spot"), r.get("max_pain")
    return {
        "expiry_close": _num(close),
        "ret_pct": _num(100 * (close / spot - 1)) if spot else None,
        "inside_range": between(r.get("range_low"), r.get("range_high")),
        "put_wall_held": None if r.get("put_wall") is None else bool(close >= r["put_wall"]),
        "call_wall_held": None if r.get("call_wall") is None else bool(close <= r["call_wall"]),
        "max_pain_closer": (None if mp is None or not spot
                            else bool(abs(close - mp) < abs(close - spot))),
        "scored_on": today,
    }


# ---------------------------------------------------------------- workbook

def _load(path: Path):
    from openpyxl import Workbook, load_workbook
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _rows(ws) -> list[dict]:
    vals = list(ws.iter_rows(values_only=True))
    if not vals:
        return []
    head = list(vals[0])
    return [dict(zip(head, r)) for r in vals[1:] if any(x is not None for x in r)]


def _date(x) -> dt.date | None:
    if isinstance(x, dt.datetime):
        return x.date()
    return x


def _write_tab(wb, name: str, rows: list[dict], cols: list[str], pos: int | None = None):
    from openpyxl.styles import Font
    if name in wb.sheetnames:
        pos = wb.sheetnames.index(name)
        wb.remove(wb[name])
    ws = wb.create_sheet(name, pos)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2" if name == SUMMARY else "F2"
    for r in rows:
        ws.append([r.get(c) for c in cols])
    for i, col in enumerate(cols, 1):
        width = max([len(col)] + [len(str(r.get(col) or "")) for r in rows])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width + 2, 24)
    return ws


def record(rows: list[dict], expiry: dt.date, path: Path = HISTORY,
           force: bool = False) -> int:
    """Write one cycle's tab. 0 if that cycle is already recorded."""
    wb = _load(path)
    name = sheet_name(expiry)
    if name in wb.sheetnames and not force:
        return 0
    _write_tab(wb, name, rows, COLUMNS)      # replaced in place, or appended
    _write_summary(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows)


def pending(path: Path = HISTORY, today: dt.date | None = None) -> list[dt.date]:
    """Expiries already passed whose tab has no outcome yet."""
    today = today or dt.date.today()
    if not path.exists():
        return []
    wb = _load(path)
    out = []
    for name in wb.sheetnames:
        if name == SUMMARY:
            continue
        rows = _rows(wb[name])
        if not rows:
            continue
        exp = _date(rows[0]["expiry"])
        if exp is not None and exp <= today and any(r.get("expiry_close") is None for r in rows):
            out.append(exp)
    return sorted(out)


def score(expiry: dt.date, closes: dict[str, float], path: Path = HISTORY,
          today: dt.date | None = None) -> int:
    """Fill in a cycle's outcome columns. Returns rows scored."""
    today = today or dt.date.today()
    wb = _load(path)
    name = sheet_name(expiry)
    rows = _rows(wb[name])
    n = 0
    for r in rows:
        if r.get("expiry_close") is None and r["symbol"] in closes:
            r.update(score_row(r, closes[r["symbol"]], today))
            n += 1
    _write_tab(wb, name, rows, COLUMNS)
    _write_summary(wb)
    wb.save(path)
    return n


def _rate(rows: list[dict], col: str):
    vals = [r[col] for r in rows if r.get(col) is not None]
    return _num(100 * sum(bool(v) for v in vals) / len(vals), 1) if vals else None


def _up_gap(rows: list[dict], col: str, top_share: float) -> float | None:
    """Points: % that rose among the top `top_share` of rows by `col`, minus %
    among the bottom `top_share` (fifths) or among the remainder (shares
    above a fifth, used for the four indices)."""
    rs = [r for r in rows if r.get(col) is not None and r.get("ret_pct") is not None]
    if len(rs) < 4:
        return None
    rs.sort(key=lambda r: r[col])
    n = len(rs)
    k = max(1, round(n * top_share))
    top = rs[-k:]
    bottom = rs[:k] if top_share <= 0.2 else rs[:n - k]
    up = lambda xs: 100 * sum(r["ret_pct"] > 0 for r in xs) / len(xs)
    return _num(up(top) - up(bottom), 1)


# The two directional hints the backtest found but did not confirm
# (docs/f_o/validation.md). Kept out of the live view; scored here, on
# records written before the outcome, so the forward record decides.
HINTS = {"equity": ("PCR: top fifth minus bottom fifth, % up (backtest: +5 to +6)",
                    "pcr_oi", 0.2),
         "index": ("skew: top 40% minus rest, % up (backtest: about +17)",
                   "skew_pts", 0.4)}


def summary_rows(cycles: dict[str, list[dict]]) -> list[dict]:
    """Per expiry and kind, plus an all-cycles total: actual vs expected."""
    out = []

    def line(label, kind, rows, gap=None):
        scored = [r for r in rows if r.get("expiry_close") is not None]
        exp = [r["expected_hit_pct"] for r in rows if r.get("expected_hit_pct") is not None]
        hint, col, share = HINTS[kind]
        if gap is None and scored:
            gap = _up_gap(scored, col, share)
        out.append({
            "cycle": label, "kind": kind, "recorded": len(rows), "scored": len(scored),
            "inside_range_pct": _rate(scored, "inside_range"),
            "expected_hit_pct": _num(sum(exp) / len(exp), 1) if exp else None,
            "put_wall_held_pct": _rate(scored, "put_wall_held"),
            "call_wall_held_pct": _rate(scored, "call_wall_held"),
            "max_pain_closer_pct": _rate(scored, "max_pain_closer"),
            "tentative_hint": hint, "hint_gap_pts": gap,
        })

    every: list[dict] = []
    for name, rows in cycles.items():
        every += rows
        for kind in ("index", "equity"):
            sub = [r for r in rows if r.get("kind") == kind]
            if sub:
                line(name, kind, sub)
    for kind in ("index", "equity"):
        sub = [r for r in every if r.get("kind") == kind]
        if not sub:
            continue
        if kind == "equity":
            # Within-cycle gaps averaged: a cycle is one recording date, so
            # this is the backtest's within-date comparison. Pooling across
            # cycles would let one cycle's market move decide it.
            gaps = [g for g in (_up_gap([r for r in rows if r.get("kind") == kind
                                         and r.get("expiry_close") is not None],
                                        "pcr_oi", 0.2) for rows in cycles.values())
                    if g is not None]
            line("ALL CYCLES", kind, sub, _num(sum(gaps) / len(gaps), 1) if gaps else None)
        else:
            line("ALL CYCLES", kind, sub)
    return out


SUMMARY_COLS = ["cycle", "kind", "recorded", "scored", "inside_range_pct",
                "expected_hit_pct", "put_wall_held_pct", "call_wall_held_pct",
                "max_pain_closer_pct", "tentative_hint", "hint_gap_pts"]


def _write_summary(wb) -> None:
    cycles = {s: _rows(wb[s]) for s in wb.sheetnames if s != SUMMARY}
    ws = _write_tab(wb, SUMMARY, summary_rows(cycles), SUMMARY_COLS, 0)
    ws.cell(ws.max_row + 2, 1,
            "One expiry is noise: stocks move together, so one large market move "
            "can push most of a cycle outside its range. Judge inside_range_pct "
            "against expected_hit_pct over several cycles. Walls held/max pain "
            "closer are positioning checks -- the backtest found no edge in them. "
            "hint_gap_pts scores the two unconfirmed directional hints; they stay "
            "out of the live view until many cycles here agree with the backtest.")


# ------------------------------------------------------------------- live

def default_symbols() -> list[str]:
    """The four main indices plus every NIFTY 50 stock that has F&O."""
    from data.sources import nse_client as nse
    from data.sources import nse_derivatives as nsed
    fno = set(nsed.fno_underlyings()["symbol"])
    cache = Path(__file__).parent / "data" / "cache" / "nifty50.csv"
    n50 = nse.fetch_csv(nse.NIFTY50_LIST_URL, cache)["Symbol"].str.strip()
    return [s for s in INDICES if s in fno] + sorted(s for s in n50 if s in fno)


def expiry_closes(expiry: dt.date) -> dict[str, float]:
    """Each underlying's close on expiry day, from that day's F&O bhavcopy."""
    from data import fo_bhavcopy as fb
    norm = fb.fetch_day(expiry)
    if norm is None or norm.empty:      # failed, or not published yet: retry later
        return {}
    sp = fb.spots(norm, expiry)
    return dict(zip(sp["symbol"], sp["spot"].astype(float)))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--symbols", nargs="+", default=None,
                    help="default: NIFTY, BANKNIFTY, FINNIFTY, MIDCPNIFTY + NIFTY 50 F&O stocks")
    ap.add_argument("--force", action="store_true",
                    help="replace this cycle's record instead of keeping it")
    ap.add_argument("--check-only", action="store_true",
                    help="only score expired cycles; record nothing")
    ap.add_argument("--path", type=Path, default=HISTORY)
    a = ap.parse_args(argv)
    today = dt.date.today()

    for exp in pending(a.path, today):
        closes = expiry_closes(exp)
        if not closes:
            print(f"{sheet_name(exp)}: expiry-day bhavcopy not published yet; "
                  f"will score on a later run.")
            continue
        n = score(exp, closes, a.path, today)
        print(f"{sheet_name(exp)}: scored {n} underlyings.")
    if a.check_only:
        return 0

    import fno_report
    from data.sources import nse_derivatives as nsed
    symbols = [s.upper() for s in (a.symbols or default_symbols())]
    rows, expiry = [], None
    for sym in symbols:
        try:
            res = fno_report.load(sym)
        except (nsed.NotFnO, ValueError, RuntimeError) as exc:
            print(f"  {sym}: skipped ({exc})", file=sys.stderr)
            continue
        expiry = expiry or res.view.expiry
        if res.view.expiry != expiry:     # an underlying on a different monthly date
            print(f"  {sym}: skipped (expiry {res.view.expiry}, cycle is {expiry})",
                  file=sys.stderr)
            continue
        rows.append(record_row(res.view, today))
    if not rows:
        print("Nothing recorded: no chain could be read.")
        return 1
    n = record(rows, expiry, a.path, force=a.force)
    if n == 0:
        print(f"'{sheet_name(expiry)}' already recorded in {a.path} -- the first "
              f"record of a cycle is kept. Use --force to replace it.")
    else:
        print(f"Recorded {n} underlyings in '{sheet_name(expiry)}' -> {a.path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

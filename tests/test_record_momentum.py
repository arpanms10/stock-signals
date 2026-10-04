"""Monthly momentum record: one tab per month, never silently rewritten."""
import datetime as dt

from openpyxl import load_workbook

import record_momentum as rm


def pick(sym, rank=1, price=100.0):
    return {"symbol": sym, "status": "Both", "rank_main": rank,
            "score_main": 250.0, "rank_cmp": rank, "score_cmp": 300.0,
            "price": price, "ret_main": 50.0, "ret_cmp": 60.0, "ret_1m": 2.0,
            "vol": 30.0, "quality": 72.0, "sector": "IT", "held": True}


def meta(sheet, day):
    return {"sheet": sheet, "recorded_on": day, "price_date": day,
            "benchmark": "NIFTY 500", "benchmark_close": 21000.0}


def read(path, sheet):
    ws = load_workbook(path)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_tab_is_named_for_the_month_with_its_year():
    assert rm.sheet_name(dt.date(2026, 10, 2)) == "Oct 2026"
    assert rm.sheet_name(dt.date(2027, 10, 1)) == "Oct 2027"


def test_first_snapshot_of_the_month_is_kept(tmp_path):
    p = tmp_path / "picks.xlsx"
    oct2 = dt.date(2026, 10, 2)
    assert rm.record([pick("A"), pick("B", 2)], meta("Oct 2026", oct2), p) == 2
    # A later run the same month must not overwrite what was suggested.
    assert rm.record([pick("C")], meta("Oct 2026", oct2), p) == 0
    assert [r["symbol"] for r in read(p, "Oct 2026")] == ["A", "B"]


def test_new_month_is_a_new_tab_and_force_replaces_in_place(tmp_path):
    p = tmp_path / "picks.xlsx"
    rm.record([pick("A")], meta("Oct 2026", dt.date(2026, 10, 2)), p)
    rm.record([pick("X")], meta("Nov 2026", dt.date(2026, 11, 2)), p)
    assert rm.record([pick("C")], meta("Oct 2026", dt.date(2026, 10, 2)), p,
                     force=True) == 1
    assert load_workbook(p).sheetnames == ["Oct 2026", "Nov 2026"]
    assert [r["symbol"] for r in read(p, "Oct 2026")] == ["C"]
    assert rm.recorded_months(p) == {"Oct 2026", "Nov 2026"}


def test_values_are_real_numbers_and_holdings_are_not_written(tmp_path):
    """Numbers stay numeric so Excel can calculate with them; the file lives
    in git, and what you hold does not."""
    p = tmp_path / "picks.xlsx"
    rm.record([pick("A")], meta("Oct 2026", dt.date(2026, 10, 2)), p)
    row = read(p, "Oct 2026")[0]
    assert "held" not in row and "month" not in row
    assert row["price"] == 100.0 and row["benchmark_close"] == 21000.0
    assert row["rank_main"] == 1 and isinstance(row["price_date"], dt.datetime)

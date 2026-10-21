"""Workbook or CSV: the files you fill in read and write the same either way."""
import openpyxl
import pytest

import portfolio as pf
import userfiles as uf
import watching as wg
from mf import holdings as mfh


def test_template_has_headers_only_and_a_guide(tmp_path):
    p = uf.template(uf.HOLDINGS, tmp_path)
    assert p == tmp_path / "holdings.xlsx"
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Holdings", uf.GUIDE_SHEET]
    rows = list(wb["Holdings"].iter_rows(values_only=True))
    assert list(rows[0]) == uf.HOLDINGS.names and len(rows) == 1     # no example rows
    guide = " ".join(str(c) for r in wb[uf.GUIDE_SHEET].iter_rows(values_only=True)
                     for c in r if c)
    assert "avg_price" in guide and "gitignore" in guide
    assert uf.template(uf.HOLDINGS, tmp_path) is None                 # never overwrites


def test_template_skipped_when_csv_exists(tmp_path):
    (tmp_path / "holdings.csv").write_text("symbol,quantity\nTCS,1\n")
    assert uf.template(uf.HOLDINGS, tmp_path) is None


def test_holdings_read_from_workbook_like_csv(tmp_path):
    p = tmp_path / "holdings.xlsx"
    uf.write(p, uf.HOLDINGS, [{"symbol": "tcs", "quantity": "10", "avg_price": "3,100",
                               "lt_quantity": "6", "bucket": "Core"}])
    got = pf.read_csv_holdings(p)
    assert got[0]["tradingsymbol"] == "TCS" and got[0]["quantity"] == 10
    assert got[0]["average_price"] == 3100 and got[0]["st_quantity"] == 4
    assert got[0]["bucket"] == "core"


def test_blank_tax_split_is_unknown_not_zero(tmp_path):
    p = tmp_path / "holdings.xlsx"
    uf.write(p, uf.HOLDINGS, [{"symbol": "TCS", "quantity": "10"}])
    h = pf.read_csv_holdings(p)[0]
    assert h["lt_quantity"] is None and h["st_quantity"] is None


def test_rewriting_a_workbook_keeps_the_guide(tmp_path):
    p = uf.template(uf.WATCHING, tmp_path)
    wg.save([wg.Watched("INFY", source="manual")], p)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames[0] == "Watching" and uf.GUIDE_SHEET in wb.sheetnames
    assert [w.symbol for w in wg.load(p)] == ["INFY"]


def test_watching_csv_still_works(tmp_path):
    p = tmp_path / "watching.csv"
    wg.save([wg.Watched("TCS", baseline_price=100.0)], p)
    assert p.read_text().startswith("#")
    assert wg.load(p)[0].baseline_price == 100.0


def test_workbook_wins_and_conflict_is_reported(tmp_path):
    (tmp_path / "watching.csv").write_text("symbol\nTCS\n")
    uf.template(uf.HOLDINGS, tmp_path)
    assert uf.find("watching", tmp_path).suffix == ".csv"
    uf.write(tmp_path / "watching.xlsx", uf.WATCHING, [])
    assert uf.find("watching", tmp_path).suffix == ".xlsx"
    assert any("watching" in w for w in uf.conflicts(tmp_path))


def test_mf_workbook_reads_the_data_sheet_not_the_guide(tmp_path):
    p = tmp_path / "mf_holdings.xlsx"
    uf.write(p, uf.MF_HOLDINGS, [{"fund": "Parag Parikh Flexi Cap", "pct": "100"}])
    wb = openpyxl.load_workbook(p)
    wb.move_sheet(uf.GUIDE_SHEET, offset=-1)          # guide first, by hand
    wb.save(p)
    raw = mfh.read_raw(p)
    assert list(raw["fund"]) == ["Parag Parikh Flexi Cap"]


def test_rows_starting_with_hash_are_comments(tmp_path):
    p = tmp_path / "holdings.xlsx"
    uf.write(p, uf.HOLDINGS, [{"symbol": "# example: RELIANCE", "quantity": "1"},
                              {"symbol": "TCS", "quantity": "2"}])
    assert [h["tradingsymbol"] for h in pf.read_csv_holdings(p)] == ["TCS"]

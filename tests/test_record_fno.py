"""F&O forward log: one tab per expiry, kept once, scored after expiry."""
import datetime as dt

from openpyxl import load_workbook

import record_fno as rf

EXP = dt.date(2026, 10, 27)


def row(sym, kind="equity", spot=100.0, lo=94.0, hi=106.0, put=95.0, call=110.0,
        mp=104.0, hit=61):
    return {"recorded_on": dt.date(2026, 10, 10), "expiry": EXP, "symbol": sym,
            "kind": kind, "spot": spot, "range_low": lo, "range_high": hi,
            "put_wall": put, "call_wall": call, "max_pain": mp,
            "expected_hit_pct": hit}


def test_first_record_of_a_cycle_is_kept(tmp_path):
    p = tmp_path / "f.xlsx"
    assert rf.record([row("A"), row("B")], EXP, p) == 2
    assert rf.record([row("C")], EXP, p) == 0
    assert rf.record([row("C")], EXP, p, force=True) == 1
    wb = load_workbook(p)
    assert wb.sheetnames == ["Summary", "Oct 2026 expiry"]


def test_score_row():
    s = rf.score_row(row("A"), 103.0, EXP)
    assert s["inside_range"] is True and s["put_wall_held"] is True
    assert s["call_wall_held"] is True and s["max_pain_closer"] is True
    assert s["ret_pct"] == 3.0
    s = rf.score_row(row("A"), 92.0, EXP)
    assert s["inside_range"] is False and s["put_wall_held"] is False
    assert s["max_pain_closer"] is False       # 92 is nearer 100 than 104


def test_pending_then_scored_and_summarised(tmp_path):
    p = tmp_path / "f.xlsx"
    rf.record([row("A"), row("B"), row("NIFTY", kind="index", hit=49)], EXP, p)
    assert rf.pending(p, EXP - dt.timedelta(days=1)) == []      # not expired
    assert rf.pending(p, EXP) == [EXP]
    # B has no close yet (e.g. delisted from F&O): stays pending
    assert rf.score(EXP, {"A": 103.0, "NIFTY": 120.0}, p, EXP) == 2
    assert rf.pending(p, EXP) == [EXP]
    assert rf.score(EXP, {"B": 90.0}, p, EXP) == 1
    assert rf.pending(p, EXP) == []

    ws = load_workbook(p)["Summary"]
    summ = [dict(zip([c.value for c in ws[1]], [c.value for c in r]))
            for r in ws.iter_rows(min_row=2) if r[1].value]
    eq = next(s for s in summ if s["cycle"] == "ALL CYCLES" and s["kind"] == "equity")
    assert eq["scored"] == 2 and eq["inside_range_pct"] == 50.0
    assert eq["expected_hit_pct"] == 61.0
    ix = next(s for s in summ if s["cycle"] == "ALL CYCLES" and s["kind"] == "index")
    assert ix["inside_range_pct"] == 0.0


def test_new_cycle_goes_after_existing_ones(tmp_path):
    p = tmp_path / "f.xlsx"
    rf.record([row("A")], EXP, p)
    nov = [dict(row("A"), expiry=dt.date(2026, 11, 23))]
    rf.record(nov, dt.date(2026, 11, 23), p)
    assert load_workbook(p).sheetnames == ["Summary", "Oct 2026 expiry", "Nov 2026 expiry"]

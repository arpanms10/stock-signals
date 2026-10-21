"""The files you fill in: holdings, mutual fund holdings, the watch list.

Each can be an Excel workbook or a CSV in config/:

  config/holdings.xlsx      or  config/holdings.csv
  config/mf_holdings.xlsx   or  config/mf_holdings.csv
  config/watching.xlsx      or  config/watching.csv

A workbook has the data on its first sheet and a "How to fill" sheet beside
it, written by setup_project.py, saying what each column means and where to
find it. A CSV carries the same help as # comment lines at the top. When both
exist, the workbook is read and a warning says so.

Everything that reads or writes these files goes through here, so the app
behaves the same whichever format you keep. Writing a workbook replaces only
the data sheet: the instructions stay.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).parent
CONFIG = ROOT / "config"
GUIDE_SHEET = "How to fill"


@dataclass(frozen=True)
class Column:
    name: str
    required: bool
    meaning: str
    example: str = ""
    where: str = ""


@dataclass(frozen=True)
class Spec:
    stem: str                 # file name without extension
    sheet: str                # the data sheet's name in a workbook
    title: str
    intro: list[str]
    columns: list[Column]
    notes: list[str] = field(default_factory=list)

    @property
    def names(self) -> list[str]:
        return [c.name for c in self.columns]


PRIVACY = ("Stays on your computer. config/ files with your positions are listed in "
           ".gitignore, so they are never uploaded to GitHub, even if you share the "
           "project.")

HOLDINGS = Spec(
    "holdings", "Holdings", "Your share holdings",
    ["One row per stock you own. Fill this sheet, save, and the dashboard's Stocks "
     "page and Today page use it on their next load.",
     "Easiest route: download your holdings from your broker (Zerodha Console, "
     "Kite, Groww ...) and run  import_kite_holdings.py  -- it fills this file "
     "for you. Otherwise type the rows in by hand.",
     "Only symbol and quantity are required. Everything else makes the advice "
     "better."],
    [Column("symbol", True, "The stock's NSE symbol, in capitals.", "RELIANCE",
            "Your broker's holdings page, or nseindia.com"),
     Column("quantity", True, "Number of shares you hold.", "40"),
     Column("lt_quantity", False, "Shares held for more than 12 months (long-term "
            "for tax). Give this or st_quantity; the other is worked out.", "40",
            "Broker holdings statement (Console: 'Long term quantity')"),
     Column("st_quantity", False, "Shares held for 12 months or less (short-term).",
            "0", "Broker holdings statement"),
     Column("avg_price", False, "Your average buying price per share, in rupees. "
            "Makes profit, loss and stop levels real.", "831.76",
            "Broker holdings page ('Avg. cost')"),
     Column("purchase_date", False, "When you bought, as YYYY-MM-DD. Leave blank if "
            "you bought over time.", "2023-04-11", "Your contract notes"),
     Column("bucket", False, "core = owned for a decade or more; satellite = one "
            "to ten years; trade = bought to trade (the only kind sold on momentum "
            "rank). Leave blank and the app suggests core or satellite.", "core"),
     Column("notes", False, "Anything for yourself.", "bought on results dip")],
    [PRIVACY,
     "Use the NSE symbol, not the company name: 'HDFCBANK', not 'HDFC Bank Ltd'.",
     "Bonds, REITs, ETFs and gold bonds from a broker file are skipped on import; "
     "the app covers company shares."])

MF_HOLDINGS = Spec(
    "mf_holdings", "MF holdings", "Your mutual fund book, as percentages",
    ["One row per fund. Give each fund's share of your whole mutual fund "
     "portfolio -- percentages only, never amounts.",
     "The Mutual funds page compares each fund with others in its category and "
     "says hold, add, trim or exit."],
    [Column("fund", True, "The fund's name as you would say it, or its AMFI scheme "
            "code. Direct or regular plan doesn't matter.", "Parag Parikh Flexi Cap",
            "Your CAMS / KFintech statement, or the fund house's site"),
     Column("pct", True, "Share of your whole MF portfolio, in %. The column "
            "should add up to 100.", "30", "Value of this fund / total MF value x 100"),
     Column("held_over_1y", False, "yes / no / blank. Only changes the exit-load "
            "and short-term tax warning on a sell verdict.", "yes"),
     Column("notes", False, "Anything for yourself.", "")],
    [PRIVACY,
     "If a name matches several funds, the page lists the candidates; use the "
     "AMFI code to be exact."])

WATCHING = Spec(
    "watching", "Watching", "Stocks you are watching closely",
    ["Names you want to keep an eye on before trusting them with money. The app "
     "adds to this list itself when you press 'Watch' on a suggestion, and dates "
     "every stop and target it reaches. You can also add rows by hand: symbol is "
     "enough, the app fills the price and plan on its next run."],
    [Column("symbol", True, "NSE symbol.", "TCS"),
     Column("added_on", False, "When you started watching (YYYY-MM-DD).", "2026-10-01"),
     Column("baseline_date", False, "When the app first flagged it. Filled by the app.", ""),
     Column("baseline_price", False, "Price on that date. Filled by the app.", ""),
     Column("stop", False, "Stop-loss level. Filled by the app.", ""),
     Column("t1", False, "First target. Filled by the app.", ""),
     Column("t2", False, "Second target. Filled by the app.", ""),
     Column("stopped_on", False, "When the stop was reached. Filled by the app.", ""),
     Column("t1_hit_on", False, "When T1 was reached. Filled by the app.", ""),
     Column("t2_hit_on", False, "When T2 was reached. Filled by the app.", ""),
     Column("source", False, "Why it is here: rank | advisor | swing | manual.", "manual"),
     Column("conviction", False, "Your own note: watch | interested | avoid.", "watch"),
     Column("notes", False, "Anything for yourself.", "")],
    ["Columns marked 'Filled by the app' can be left blank."])

SPECS = {s.stem: s for s in (HOLDINGS, MF_HOLDINGS, WATCHING)}


# ------------------------------------------------------------------ paths

def find(stem: str, config: Path = CONFIG) -> Path | None:
    """The file in use for `stem`: the workbook if there is one, else the CSV."""
    x, c = config / f"{stem}.xlsx", config / f"{stem}.csv"
    return x if x.exists() else c if c.exists() else None


def path_for(stem: str, config: Path = CONFIG) -> Path:
    """Where to read or write `stem`: the existing file, else the CSV."""
    return find(stem, config) or config / f"{stem}.csv"


def conflicts(config: Path = CONFIG) -> list[str]:
    """Warnings for files kept in both formats (the workbook wins)."""
    return [f"Both {s}.xlsx and {s}.csv are in config/; the .xlsx is used. "
            f"Delete the one you no longer keep."
            for s in SPECS if (config / f"{s}.xlsx").exists()
            and (config / f"{s}.csv").exists()]


# ---------------------------------------------------------------- reading

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return str(v).strip()


def data_sheet(wb):
    """The first sheet that is not the instructions."""
    for name in wb.sheetnames:
        if name.strip().lower() != GUIDE_SHEET.lower():
            return wb[name]
    return wb[wb.sheetnames[0]]


def read(path: Path | str) -> tuple[list[str], list[dict]]:
    """(header, rows) from a workbook's data sheet or a CSV. Lines or rows
    starting with # are skipped; values come back as stripped strings."""
    path = Path(path)
    if not path.exists():
        return [], []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = data_sheet(wb)
        ws.reset_dimensions()
        grid = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        grid = [r for r in grid if any(r) and not r[0].startswith("#")]
        if not grid:
            return [], []
        head = [h for h in grid[0]]
        while head and not head[-1]:
            head.pop()
        rows = [dict(zip(head, r[:len(head)])) for r in grid[1:]]
        return head, rows
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(l for l in fh if not l.lstrip().startswith("#"))
        rows = [{k: (v or "").strip() for k, v in r.items() if k is not None}
                for r in reader]
        return list(reader.fieldnames or []), rows


# ---------------------------------------------------------------- writing

def csv_help(spec: Spec) -> str:
    width = max(len(c.name) for c in spec.columns) + 2
    lines = [f"# {spec.title}. Opens in Excel.", "#"]
    for c in spec.columns:
        tag = "required. " if c.required else ""
        lines.append(f"#   {c.name.ljust(width)}{tag}{c.meaning}")
    return "\n".join(lines) + "\n"


def _guide(wb, spec: Spec) -> None:
    from openpyxl.styles import Alignment, Font
    if GUIDE_SHEET in wb.sheetnames:
        del wb[GUIDE_SHEET]
    ws = wb.create_sheet(GUIDE_SHEET)
    ws["A1"] = spec.title
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for line in spec.intro:
        ws.cell(r, 1, line)
        r += 1
    r += 1
    head = ["Column", "Required?", "What it means", "Example", "Where to find it"]
    for i, h in enumerate(head, 1):
        ws.cell(r, i, h).font = Font(bold=True)
    for c in spec.columns:
        r += 1
        for i, v in enumerate((c.name, "yes" if c.required else "optional",
                               c.meaning, c.example, c.where), 1):
            ws.cell(r, i, v).alignment = Alignment(wrap_text=True, vertical="top")
    if spec.notes:
        r += 2
        ws.cell(r, 1, "Good to know").font = Font(bold=True)
        for n in spec.notes:
            r += 1
            ws.cell(r, 1, n)
    for col, w in zip("ABCDE", (16, 11, 60, 22, 40)):
        ws.column_dimensions[col].width = w


def write(path: Path | str, spec: Spec, rows: list[dict],
          columns: list[str] | None = None) -> None:
    """Write rows under `columns` (default: the spec's). A workbook keeps or
    gains its "How to fill" sheet; a CSV gets the help as comments."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = columns or spec.names
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font
        if path.exists():
            wb = load_workbook(path)
            old = data_sheet(wb)
            title, pos = old.title, wb.sheetnames.index(old.title)
            if old.title.lower() == GUIDE_SHEET.lower():
                title, pos = spec.sheet, 0
            else:
                del wb[old.title]
            ws = wb.create_sheet(title, pos)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = spec.sheet
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
        for i, c in enumerate(columns, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(c) + 4)
        if GUIDE_SHEET not in wb.sheetnames:
            _guide(wb, spec)
        wb.active = wb.sheetnames.index(ws.title)
        wb.save(path)
        return
    with path.open("w", newline="", encoding="utf-8") as fh:
        fh.write(csv_help(spec))
        w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in columns})


def template(spec: Spec, config: Path = CONFIG) -> Path | None:
    """A blank workbook for `spec`, unless the user already has the file in
    either format. Returns the path written, or None."""
    if find(spec.stem, config) is not None:
        return None
    path = config / f"{spec.stem}.xlsx"
    write(path, spec, [])
    return path

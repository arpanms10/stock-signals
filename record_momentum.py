"""Record this month's momentum picks, once, for checking later.

A backtest says what the ranking WOULD have done; this file says what it
actually suggested, on the day, at the price on the day. Months from now that
is the only honest way to measure how far the live picks drifted from the
backtest -- the ranking code will have changed, and recomputing old months
with new code measures the new code, not the advice you were given.

Writes history/momentum_picks.xlsx, one tab per month named for the PRICE
date ("Oct 2026"): every stock in the top band on either lookback (the
dashboard's Momentum tab), with its ranks, scores, price and returns, plus the
benchmark level, so a later run can compute each pick's return and its return
over the index from the same starting point.

One tab per month -- the strategy rebalances monthly, and the first record of
a month is the one that counts. Later runs that month do nothing unless
--force replaces that tab. New months are added as a new tab at the end.

What you held is deliberately left out. This file is meant to live in git,
and holdings do not.

    PYTHONPATH=. .venv/bin/python record_momentum.py
    PYTHONPATH=. .venv/bin/python record_momentum.py --force     # redo this month
"""
from __future__ import annotations

import argparse
import datetime as dt
from pathlib import Path

HISTORY = Path(__file__).parent / "history" / "momentum_picks.xlsx"

# The month is the tab name, so it is not repeated as a column.
COLUMNS = ["recorded_on", "price_date", "universe_size", "n_hold",
           "lookback_main", "lookback_cmp", "regime", "benchmark",
           "benchmark_close", "symbol", "list", "rank_main", "score_main",
           "rank_cmp", "score_cmp", "price", "ret_main_pct", "ret_cmp_pct",
           "ret_1m_pct", "vol_pct", "quality", "sector"]


def sheet_name(price_date: dt.date) -> str:
    """Tab name like "Oct 2026" -- with the year, so next October cannot clash."""
    return f"{price_date:%b %Y}"


def recorded_months(path: Path = HISTORY) -> set[str]:
    """Tab names already in the workbook."""
    if not path.exists():
        return set()
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        return set(wb.sheetnames)
    finally:
        wb.close()


def _num(v, nd: int = 2):
    """A real number for Excel, rounded; blank for missing."""
    if v is None or v != v:
        return None
    return round(float(v), nd) if nd else int(round(float(v)))


def pick_row(p: dict) -> dict:
    return {
        "symbol": p["symbol"], "list": p["status"],
        "rank_main": _num(p.get("rank_main"), 0),
        "score_main": _num(p.get("score_main"), 1),
        "rank_cmp": _num(p.get("rank_cmp"), 0),
        "score_cmp": _num(p.get("score_cmp"), 1),
        "price": _num(p.get("price")),
        "ret_main_pct": _num(p.get("ret_main"), 1),
        "ret_cmp_pct": _num(p.get("ret_cmp"), 1),
        "ret_1m_pct": _num(p.get("ret_1m"), 1),
        "vol_pct": _num(p.get("vol"), 1),
        "quality": _num(p.get("quality"), 0),
        "sector": p.get("sector", ""),
    }


def write_sheet(rows: list[dict], sheet: str, path: Path = HISTORY,
                force: bool = False) -> int:
    """Write one month's tab. Returns rows written; 0 if already recorded."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    pos = len(wb.sheetnames)
    if sheet in wb.sheetnames:
        if not force:
            return 0
        pos = wb.sheetnames.index(sheet)
        wb.remove(wb[sheet])
    ws = wb.create_sheet(sheet, pos)
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([r.get(c) for c in COLUMNS])
    for i, col in enumerate(COLUMNS, 1):
        width = max(len(col), *(len(str(r.get(col) or "")) for r in rows))
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width + 2, 34)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows)


def record(picks: list[dict], meta: dict, path: Path = HISTORY,
           force: bool = False) -> int:
    """One month's picks into its own tab. `meta` holds the month-level
    fields, plus "sheet" (the tab name)."""
    base = {c: meta.get(c) for c in COLUMNS}
    rows = [{**base, **pick_row(p)} for p in picks]
    return write_sheet(rows, meta["sheet"], path, force)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--universe-size", type=int, default=500,
                    help="most liquid stocks to rank (the dashboard default)")
    ap.add_argument("--force", action="store_true",
                    help="replace this month's record instead of keeping it")
    ap.add_argument("--path", type=Path, default=HISTORY)
    args = ap.parse_args()

    from data import store
    from scoring import timing
    from ui import service

    snap = service.build(universe_size=args.universe_size)
    sheet = sheet_name(snap.as_of)
    if sheet in recorded_months(args.path) and not args.force:
        print(f"'{sheet}' already recorded in {args.path} -- the first snapshot "
              f"of the month is kept. Use --force to replace it.")
        return
    if len(snap.lookbacks) < 2 or not snap.momentum_table:
        raise SystemExit("No momentum ranking to record -- refresh market data "
                         "first.")

    cfg = timing.load_config()
    bench_name = cfg["signals"]["regime_index"]
    bench = store.load_index(store.connect(), bench_name)
    bench = bench[bench["date"] <= snap.as_of] if not bench.empty else bench
    meta = {
        "sheet": sheet, "recorded_on": dt.date.today(),
        "price_date": snap.as_of,
        "universe_size": args.universe_size,
        "n_hold": cfg["momentum_strategy"]["n_hold"],
        "lookback_main": snap.lookbacks[0], "lookback_cmp": snap.lookbacks[1],
        "regime": snap.regime, "benchmark": bench_name,
        "benchmark_close": _num(float(bench["close"].iloc[-1]))
        if not bench.empty else None,
    }
    n = record(snap.momentum_table, meta, args.path, force=args.force)
    print(f"Recorded {n} picks in tab '{sheet}' (prices of {snap.as_of}) "
          f"-> {args.path}")


if __name__ == "__main__":
    main()

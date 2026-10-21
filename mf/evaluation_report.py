"""run_mf.py --excel: verdicts, side-by-side comparison, and each category's table."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from mf import engine, verdict
from mf.backtest import calibration_line
from mf.report import PCT1, _lines, _write_table

ACTION_COLOURS = {"EXIT": "C0392B", "AVOID": "C0392B", "TRIM": "E67E22",
                  "WATCH": "F1C40F", "ADD": "27AE60", "HOLD": "7F8C8D",
                  "NEUTRAL": "95A5A6"}

METRIC_COLS = {
    "cagr_1y": "CAGR 1y", "cagr_3y": "CAGR 3y", "cagr_5y": "CAGR 5y",
    "excess_1y": "vs median 1y", "excess_3y": "vs median 3y", "excess_5y": "vs median 5y",
    "rank_1y": "Rank 1y", "rank_3y": "Rank 3y", "rank_5y": "Rank 5y",
    "vs_index_3y": "vs index fund 3y", "consistency_3y": "Beat median (rolling 1y)",
    "vol_3y": "Volatility 3y", "max_drawdown_3y": "Max drawdown 3y",
    "downside_capture_3y": "Downside capture 3y", "history_months": "Months of history",
}


def _colour_actions(ws, col: int, first_row: int, last_row: int) -> None:
    for r in range(first_row, last_row + 1):
        c = ws.cell(row=r, column=col)
        colour = ACTION_COLOURS.get(str(c.value))
        if colour:
            c.fill = PatternFill("solid", fgColor=colour)
            c.font = Font(bold=True, color="FFFFFF")


def build(path: Path, ev: engine.Evaluation) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rules = verdict.load_rules()
    wb = Workbook()

    ws = wb.active
    ws.title = "Verdicts"
    rows = [{
        "Fund": d.name, "Category": d.category, "Held %": d.weight,
        "Action": d.action, "Target %": d.target, "State": d.state,
        "Why": "; ".join(d.reasons), "Cautions": "; ".join(d.cautions),
        "Same-category alternatives": "; ".join(d.alternatives),
    } for d in ev.decisions]
    df = pd.DataFrame(rows)
    end = _write_table(ws, df, formats={"Held %": "0.0", "Target %": "0.0"},
                       widths={"Fund": 44, "Category": 22, "Why": 70, "Cautions": 50,
                               "Same-category alternatives": 60},
                       wrap=("Why", "Cautions", "Same-category alternatives"))
    _colour_actions(ws, 4, 2, end - 1)

    ws = wb.create_sheet("Compare")
    codes = [d.code for d in ev.decisions if d.code is not None]
    t = engine.compare_table(ev, codes).reset_index().rename(columns={"index": "Measure"})
    _write_table(ws, t, widths={"Measure": 32, **{c: 30 for c in t.columns[1:]}},
                 autofilter=False, freeze="B2")
    for r in range(2, len(t) + 2):
        label = ws.cell(row=r, column=1).value or ""
        for j in range(2, len(t.columns) + 1):
            c = ws.cell(row=r, column=j)
            if isinstance(c.value, float):
                if label.startswith(("CAGR", "vs", "Rank", "Beat", "Volatility", "Max")):
                    c.number_format = PCT1
                elif label.startswith("Downside"):
                    c.number_format = "0.00"
    _colour_actions_row = len(t) + 1
    for j in range(2, len(t.columns) + 1):
        _colour_actions(ws, j, _colour_actions_row, _colour_actions_row)

    # One sheet per category touched: every fund in it, so a verdict can be
    # checked against the whole field rather than taken on trust.
    judged = set(rules["categories"])
    for cat, g in ev.snap.groupby("category"):
        if cat not in judged:
            continue
        ws = wb.create_sheet(cat.split(":", 1)[1][:28].replace("/", "-").title())
        g = g.copy()
        g["State"] = [verdict.state(g.loc[c], rules)[0] for c in g.index]
        g["Fund"] = [ev.names.get(c, c) for c in g.index]
        g = g.sort_values("rank_3y", ascending=False, na_position="last")
        out = g[["Fund", "State"] + list(METRIC_COLS)].rename(columns=METRIC_COLS)
        fmts = {v: PCT1 for k, v in METRIC_COLS.items()
                if k not in ("downside_capture_3y", "history_months")}
        fmts["Downside capture 3y"] = "0.00"
        _write_table(ws, out, formats=fmts, widths={"Fund": 50}, freeze="C2")
        mine = {d.code for d in ev.decisions}
        for i, c in enumerate(g.index, 2):
            if c in mine:
                ws.cell(row=i, column=1).font = Font(bold=True)

    ws = wb.create_sheet("Read me")
    _lines(ws, [
        ("title", f"Mutual fund evaluation, as of {ev.as_of}"),
        ("text", "Every fund is judged against its SEBI category, on month-end NAVs of the "
                 "Direct Growth plan. Ranks are the share of category peers beaten (100% = "
                 "best). 'vs median' is CAGR minus the category median CAGR."),
        ("sub", "States (config/mf_rules.yaml)"),
        ("bullet", "WEAK: bottom quarter over 3 years, still bottom half over 1, 2+ pts/yr "
                   "behind the median -- or trailing the category's index fund by 2+ pts/yr."),
        ("bullet", "SLIPPING: a decent 3-year record but a bottom-quarter last year."),
        ("bullet", "STRONG: top 40% over 3 years, beat the median in at least half the rolling "
                   "1-year windows, not falling harder than peers in down months."),
        ("bullet", "YOUNG: under 3 years of history. OK: everything else."),
        ("sub", "Actions"),
        ("bullet", "Held: WEAK -> EXIT; over a cap or a near-twin of another holding -> TRIM; "
                   "SLIPPING/YOUNG -> WATCH; STRONG with room -> ADD; else HOLD."),
        ("bullet", "Not held: STRONG -> ADD, WEAK -> AVOID, SLIPPING/YOUNG -> WATCH, OK -> NEUTRAL."),
        ("sub", "Caveat"),
        ("text", "Past rank is no guarantee of future rank. scripts/run_mf.py --backtest "
                 "measures what each state was worth on funds that existed at the time, closed "
                 "ones included. Decision support, not advice."),
        ("text", calibration_line() or "No backtest yet: scripts/run_mf.py --backtest."),
    ])
    wb.save(path)
    return path
